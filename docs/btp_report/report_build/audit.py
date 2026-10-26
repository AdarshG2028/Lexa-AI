import docx, re

d = docx.Document("B.Tech_Project_Report_Lexa.docx")
text = "\n".join(p.text for p in d.paragraphs)
print("paragraphs", len(d.paragraphs), "tables", len(d.tables), "inline_shapes", len(d.inline_shapes))
print("words", len(text.split()))
print("asterisks", text.count("*"), "backticks", text.count("`"), "TODO", text.count("TODO"))
cites = sorted(set(int(x) for x in re.findall(r"\[(\d{1,2})\]", text)))
print("citations", cites)
heads = [p.text for p in d.paragraphs if p.style.name.startswith("BTP Heading")]
print("headings:")
for h in heads:
    print("  ", h)
